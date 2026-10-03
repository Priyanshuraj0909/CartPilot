"""Generate report/PDF/PPT from verified local evidence; no application dependency.
Run in a separate environment with python-docx, python-pptx and reportlab installed.
"""
from pathlib import Path
from html import escape
import re
from docx import Document
from docx.shared import Inches as DI, Pt as DP
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/deliverables'
SHOTS = ROOT / 'docs/screenshots'
TITLE = 'CartPilot — A Multi-Agent E-Commerce Decision and Operations Management System'
IDENTITY = '[Student / team names] • [Institution] • [Department] • [Project guide]'
ABSTRACT = (ROOT / 'PROJECT_REPORT_OUTLINE.md').read_text().split('## 1. Abstract draft (193 words)')[1].split('## 2.')[0].strip()
SECTIONS = [
('Introduction', 'CartPilot is a decision-support layer above e-commerce inventory and catalog data. It combines pricing, replenishment, promotion and listing analysis in a single merchant-scoped interface. The prototype demonstrates explainable coordination and human-controlled changes using deterministic specialist modules. It does not replace a storefront or predict business performance.'),
('Problem statement', 'Pricing, inventory, discounts and catalog quality are interdependent. A promotion may stimulate demand while stock is constrained; a weak listing may be treated as a price problem. Independent recommendations require reconciliation before acting. CartPilot addresses this coordination problem through shared context, explicit dependencies, policy validation and an auditable approval workflow. These are motivating scenarios; the project did not quantify business losses or revenue uplift.'),
('Objectives', 'The objectives are to centralize scoped product, order, inventory and price-history signals; produce four typed, explainable specialist recommendations; coordinate goals and conflicting opportunities; require human approval and fresh validation; support bounded read-only Shopify synchronization; and provide reproducible tests and a reliable academic demonstration.'),
('Technical background', 'The implementation uses typed HTTP APIs, relational persistence, deterministic decision rules and transactional workflow control. FastAPI and Pydantic support request/response validation. SQLAlchemy provides asynchronous persistence, while Alembic versions the schema. React components present merchant-scoped data and recommendation cards. The cited primary documentation explains these technologies. This report does not claim a peer-reviewed literature survey or a competitive benchmark.'),
('Proposed system and methodology', 'Sense reads product, sales, available-stock and listing facts. Decide applies four specialists and a goal-based Master Orchestrator. Act validates a typed proposal, obtains approval, asks for separate execution confirmation, rechecks current state and records the result. Learn currently means retaining outcomes for human review. No trained model, automatic retraining, reinforcement learning or LLM runtime is implemented.'),
('Architecture', 'The browser calls a FastAPI backend. Shared merchant-scoped signals feed Pricing, Restock, Promotion and Listing modules. The Master Orchestrator returns an advisory priority plan with conflicts, synergies, dependencies, blocked candidates and failures. An independent action service controls local execution. PostgreSQL is the deployment target; isolated tests and the offline demo use SQLite. Redis is optional diagnostics. Shopify import is read-only and bounded; local changes are never written back.'),
('Specialist algorithms', 'Sales velocity is eligible sold units divided by lookback days. Available stock is max(0, physical − reserved − unavailable). Pricing uses ordered demand/stock rules, Decimal rounding, movement caps and a cost-based markup floor. Restock uses lead-time demand, ceiling-rounded safety stock, buffer and an order cap. Promotion bounds discounts by inventory safety and retained gross margin. Listing quality uses weighted factual text checks; missing attributes require verification. Exact thresholds and rounding rules appear in Appendix A.'),
('Goal-based orchestration', 'The orchestrator supports six controlled merchant goals, analyzes at most 100 active products and builds a bounded candidate queue. Balanced Growth weights inventory 0.30, revenue 0.30, margin 0.20 and catalog 0.20. Critical inventory takes precedence. Conflicts include competing price changes and promotion on constrained stock. Dependencies can require listing review before discounting. Confidence, risk and priority are distinct heuristic signals; they are not calibrated probabilities.'),
('Database and API design', 'Fourteen modeled entities cover merchants, products, inventory, orders/items, price history, agent runs, recommendations, actions, approvals, audit logs and Shopify connection/external mappings. Foreign keys and unique identities support integrity and duplicate prevention. Money uses Decimal/Numeric(10,2), and API timestamps normalize to UTC. Three Alembic revisions end at 11shopify_read_only. AgentRun exists as a model; the Agent Activity UI currently records browser-session activity. Appendix B lists actual APIs and entity responsibilities.'),
('Implementation', 'The verified stack uses Python 3.12.7, FastAPI 0.142.2, Pydantic 2.13.5, SQLAlchemy 2.1.1 and Alembic 1.20.0. The frontend uses React 18.3.1, TypeScript 5.9.3, Vite 5.4.21 and Tailwind 3.4.19. PostgreSQL 16 is the container target; isolated PostgreSQL 17.11 execution was checked. These are recorded verification versions, not claims about the latest releases. Nine frontend routes cover catalog, analysis, review, audit and integrations.'),
('Approval, execution and security', 'Analysis alone changes no product values. Action creation validates immutable typed payloads and policies. Human approval and execution confirmation are separate steps. Current price, cost, stock and source facts are checked again before execution. Transactions, atomic claims, workflow identities and audits protect retries. Price/listing execution changes local data; promotion/restock execution is simulated. There is no authenticated merchant identity. Production mode blocks action writes and Shopify routes; trusted local/private access is required.'),
('Testing and verification', 'The recorded final regression passed 560 backend tests and 76 frontend tests across nine files. TypeScript and Vite build passed. Fresh SQLite migrations, seed/reseed, stale-state checks, nine direct browser routes, responsive layouts and the approval workflow were verified. Basic PostgreSQL migrations and local execution were checked in Phase 13. Shopify transport was mocked. Test counts are evidence of exercised behavior, not prediction accuracy. Docker/Nginx execution, cloud hosting, remote CI and locking stress remain unverified.'),
('Demonstration and results', 'The Four-product Demo gives the Wireless Mouse eight available units and five units/day velocity. Coverage is 1.6 days and the restock suggestion is 67 units. Pricing proposes 999.00 → 1048.95. A declining-sales Speaker supports a 10% guarded discount to 1349.10; incomplete Headphones content requires listing review. Balanced Growth prioritizes Mouse restock, Headphones listing, Speaker promotion and Mouse price review. Separate confirmation applies the local Mouse price change and records execution_completed. A 50% proposal fails the 10% movement policy. No measured sales uplift is claimed.'),
('Deployment and operations', 'The deployment topology is an HTTPS static frontend, a protected Python service and managed PostgreSQL. Migrations run once before serving; seeding is explicit and only for disposable demo data. VITE_API_URL is a public build-time value. CORS must name exact frontend origins. Production mode preserves write restrictions. A private full-workflow demo requires provider access controls. No hosting provider, cloud account or public URL has been configured for this delivery; the local demo is independent of those services.'),
('Limitations', 'The system uses fixed heuristic rules and demand assumptions. External restock/promotion execution is simulated, Shopify writes are absent and live-store verification has not been performed. Authentication, supplier orders, autonomous learning, scientific forecasting evaluation and measured business impact are outside the implemented scope. Catalog analysis is bounded, some aggregates scale by product count, and PostgreSQL concurrency stress remains unverified.'),
('Future scope', 'Potential extensions include authenticated merchant identity, ingress controls, measured forecasting, verified supplier adapters, carefully authorized external writes, additional commerce integrations and outcome evaluation. These require separate development and validation; they are not existing capabilities.'),
('Conclusion', 'CartPilot demonstrates modular, explainable e-commerce coordination with a human-controlled execution boundary. The prototype presents structured priorities, exposes conflicting proposals, blocks unsafe changes and retains audit evidence. It is suitable for a trusted academic demonstration. Production expansion requires the identity, infrastructure and external-execution work identified above.'),
('References', 'FastAPI documentation — https://fastapi.tiangolo.com/\nSQLAlchemy asyncio — https://docs.sqlalchemy.org/en/20/orm/extensions/asyncio.html\nAlembic tutorial — https://alembic.sqlalchemy.org/en/latest/tutorial.html\nReact documentation — https://react.dev/\nShopify Admin GraphQL — https://shopify.dev/docs/api/admin-graphql/latest\nProject evidence: TEST_SUMMARY.md, FINAL_STATUS.md, phases/PHASE_13.md, phases/PHASE_14.md and docs/examples/submission-browser.json. Technology references support implementation concepts; project claims derive from local source and recorded verification.')]
SLIDES = [
('CartPilot', ['The AI Manager Sitting on Top of Your Inventory', 'Explainable e-commerce coordination with guarded execution', IDENTITY], None),
('The coordination problem', ['A discount can conflict with low stock', 'Weak listing content can resemble weak demand', 'Independent proposals need a shared operational priority'], '02_products.png'),
('Project objectives', ['Four typed, explainable specialist recommendations', 'One merchant-goal priority plan', 'Policy validation, approval and persistent audit'], '01_dashboard.png'),
('Proposed solution', ['Scoped store data → shared signals', 'Specialists → conflict and dependency resolution', 'Advisory queue → separately reviewed action'], '08_ai_manager.png'),
('System architecture', ['React / TypeScript → FastAPI → relational data', 'Shopify import: read-only; Redis: optional', 'Local price/listing edits; simulated restock/promotions'], None),
('Four specialist agents', ['Pricing: bounded Decimal price rules', 'Restock: demand + safety + buffer deficit', 'Promotion: inventory and margin-safe discounts', 'Listing: grounded quality checks and suggestions'], '05_restock_agent.png'),
('Master Orchestrator', ['Balanced Growth: critical restock first', 'Listing before inappropriate demand stimulation', 'Show selected, blocked and omitted work'], '09_priority_plan.png'),
('Approval and guardrails', ['Review and approval do not execute', 'Separate confirmation + current-state revalidation', '50% price movement fails the 10% policy'], '12_execution_confirmation.png'),
('Demonstrated behavior', ['Mouse: 8 available; 5/day; restock 67', 'Approved local price: 999 → 1048.95', 'Persistent execution audit; Shopify import mocked'], '13_audit_history.png'),
('Verification evidence', ['560 backend tests; 76 frontend tests', 'TypeScript / Vite build and 9 browser routes passed', 'Fresh migrations/seed; stale-state and policy blocks', 'Basic PostgreSQL execution checked; cloud unverified'], None),
('Limits and future work', ['Deterministic heuristics; no trained forecasting or learning', 'No production authentication; trusted private access', 'External actions simulated; Shopify read-only', 'Future: identity, forecasting and authorized adapters'], None),
('Conclusion and questions', ['Explainable coordination across four domains', 'Human control, current-state safety and audit evidence', 'Trusted academic demo; no business uplift claimed'], '08_ai_manager.png')]

def clean(s):
    return re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', s).replace('`', '').replace('**', '')

def generate():
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = DP(11)
    styles = getSampleStyleSheet()
    styles['BodyText'].leading = 15
    flow = []
    def para(text, heading=False):
        if heading:
            doc.add_heading(text, level=1)
            flow.append(Paragraph(escape(text), styles['Heading1']))
        else:
            doc.add_paragraph(text)
            flow.append(Paragraph(escape(text).replace('\n','<br/>'), styles['BodyText']))
            flow.append(Spacer(1,8))
    para(TITLE, True)
    para(IDENTITY)
    para('Submission date: [Insert required date]\nPrepared from verified project evidence dated 3 October 2026. Academic identity and institution-specific formatting require completion.')
    doc.add_page_break(); flow.append(PageBreak())
    para('Contents', True)
    para('Abstract\n' + '\n'.join(f'{i}. {h}' for i,(h,_) in enumerate(SECTIONS,1)) + '\nAppendix A: Algorithms\nAppendix B: API, database and security\nAppendix C: Demonstration screenshots')
    doc.add_page_break(); flow.append(PageBreak())
    para('Abstract', True); para(ABSTRACT)
    for i,(h,body) in enumerate(SECTIONS,1):
        para(f'{i}. {h}', True); para(body)
    for title, name in [('Appendix A: Source-matched algorithms','docs/ALGORITHMS.md'), ('Appendix B: API, database and security','docs/API_DATABASE_SECURITY.md')]:
        doc.add_page_break(); flow.append(PageBreak()); para(title,True)
        for block in (ROOT/name).read_text().split('\n\n'):
            block=clean(block.strip())
            if not block or block.startswith('```'): continue
            para(block.lstrip('# ') if block.startswith('#') else block, block.startswith('#'))
    doc.add_page_break(); flow.append(PageBreak()); para('Appendix C: Real browser demonstration evidence', True)
    captions = {'01_dashboard.png':'Merchant-scoped dashboard', '04_pricing_agent.png':'Mouse pricing recommendation; heuristic confidence', '05_restock_agent.png':'Restock demand and quantity assessment', '06_promotion_agent.png':'Inventory/margin-safe Speaker promotion', '07_listing_agent.png':'Grounded Headphones listing review', '09_priority_plan.png':'Balanced Growth priority queue', '10_cross_agent_intelligence.png':'Explicit inventory/promotion conflict', '12_execution_confirmation.png':'Separate execution confirmation', '13_audit_history.png':'Persistent workflow audit', '14_policy_block.png':'Unsafe movement policy block', '15_shopify_read_only_mock.png':'Mocked Shopify read-only import; not a live store'}
    for idx,(name,caption) in enumerate(captions.items(),1):
        para(f'Figure {idx}. {caption}')
        path=SHOTS/name
        w,h=PILImage.open(path).size
        width=min(460, 570*w/h); height=width*h/w
        doc.add_picture(str(path), width=DI(width/72))
        flow.append(Image(str(path),width=width,height=height)); flow.append(Spacer(1,12))
    doc.save(OUT/'CartPilot_Final_Report.docx')
    def footer(canvas,document):
        canvas.setFont('Helvetica',8);canvas.setFillColor(colors.HexColor('#475569'))
        canvas.drawString(42,25,'CartPilot • Verified MVP • Academic details pending')
        canvas.drawRightString(553,25,str(document.page))
    SimpleDocTemplate(str(OUT/'CartPilot_Final_Report.pdf'),topMargin=42,bottomMargin=42,rightMargin=42,leftMargin=42).build(flow,onFirstPage=footer,onLaterPages=footer)
    prs=Presentation(); prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    def textbox(slide,x,y,w,h,text,size=24,color='E2E8F0'):
        box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf=box.text_frame; tf.word_wrap=True
        for i,line in enumerate(text.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line
            p.font.size=Pt(size);p.font.color.rgb=RGBColor.from_string(color);p.font.name='Aptos'
            p.space_after=Pt(18)
        return box
    for i,(title,lines,shot) in enumerate(SLIDES,1):
        slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.background.fill.solid();slide.background.fill.fore_color.rgb=RGBColor.from_string('0F172A')
        textbox(slide,.55,.3,12.2,.8,title,34,'38BDF8')
        textbox(slide,.6,1.5,5.0 if shot else 12,5.3,'\n'.join(lines),23 if shot else 28)
        if shot:
            path=SHOTS/shot; w,h=PILImage.open(path).size; pw=min(6.5,5.3*w/h);ph=pw*h/w
            slide.shapes.add_picture(str(path), Inches(6.2+(6.5-pw)/2), Inches(1.35+(5.3-ph)/2), width=Inches(pw),height=Inches(ph))
        textbox(slide,.6,7.02,12,.3,f'CartPilot | deterministic specialists • guarded local/simulated actions                                      {i:02d} / 12',10,'94A3B8')
        slide.notes_slide.notes_text_frame.text = 'Evidence: '+ str(ROOT/'PROJECT_REPORT_OUTLINE.md') + '\n' + ' '.join(lines) + '\nConfidence and priority are heuristics. Shopify is mocked/read-only. No measured business uplift is claimed.'
    prs.save(OUT/'CartPilot_Final_Presentation.pptx')
    print('Created DOCX, PDF and 12-slide PPTX in', OUT)

if __name__ == '__main__':
    generate()
